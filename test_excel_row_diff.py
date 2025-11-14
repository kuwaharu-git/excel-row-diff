#!/usr/bin/env python3
"""
テストスクリプト - Excel Row Diff のテスト
"""

from openpyxl import Workbook
import os
import sys

def create_sample_files():
    """
    テスト用のサンプルExcelファイルを作成
    """
    print("サンプルファイルを作成中...")
    
    # ファイル1を作成
    wb1 = Workbook()
    
    # Sheet1: 両方に存在するシート（一部の行が異なる）
    ws1 = wb1.active
    ws1.title = "Sheet1"
    ws1.append(["ID", "名前", "年齢", "部署"])
    ws1.append([1, "田中太郎", 30, "営業部"])
    ws1.append([2, "佐藤花子", 25, "経理部"])
    ws1.append([3, "鈴木一郎", 35, "開発部"])
    ws1.append([4, "高橋美咲", 28, "人事部"])
    ws1.append([5, "渡辺健太", 32, "営業部"])
    
    # Sheet2: 両方に存在するシート（完全に異なる）
    ws2 = wb1.create_sheet("Sheet2")
    ws2.append(["商品ID", "商品名", "価格"])
    ws2.append([101, "ノートPC", 120000])
    ws2.append([102, "マウス", 2000])
    ws2.append([103, "キーボード", 5000])
    
    # Sheet3: ファイル1にのみ存在
    ws3 = wb1.create_sheet("OnlyInFile1")
    ws3.append(["プロジェクト", "開始日", "担当者"])
    ws3.append(["プロジェクトA", "2024-01-01", "田中"])
    ws3.append(["プロジェクトB", "2024-02-01", "佐藤"])
    
    wb1.save("sample_file1.xlsx")
    print("  ✓ sample_file1.xlsx を作成")
    
    # ファイル2を作成
    wb2 = Workbook()
    
    # Sheet1: 一部の行が異なる
    ws1 = wb2.active
    ws1.title = "Sheet1"
    ws1.append(["ID", "名前", "年齢", "部署"])
    ws1.append([1, "田中太郎", 30, "営業部"])  # 同じ
    ws1.append([2, "佐藤花子", 26, "経理部"])  # 年齢が異なる（新しい行扱い）
    ws1.append([3, "鈴木一郎", 35, "開発部"])  # 同じ
    # [4, "高橋美咲", 28, "人事部"] は存在しない
    ws1.append([5, "渡辺健太", 32, "営業部"])  # 同じ
    ws1.append([6, "山田太郎", 29, "総務部"])  # 新規
    
    # Sheet2: 完全に異なる
    ws2 = wb2.create_sheet("Sheet2")
    ws2.append(["商品ID", "商品名", "価格"])
    ws2.append([201, "モニター", 30000])
    ws2.append([202, "ヘッドセット", 8000])
    
    # Sheet4: ファイル2にのみ存在
    ws4 = wb2.create_sheet("OnlyInFile2")
    ws4.append(["顧客ID", "顧客名", "契約日"])
    ws4.append([1001, "ABC商事", "2024-03-01"])
    ws4.append([1002, "XYZ株式会社", "2024-03-15"])
    
    wb2.save("sample_file2.xlsx")
    print("  ✓ sample_file2.xlsx を作成")
    
    print("\n作成されたサンプルファイル:")
    print("  - sample_file1.xlsx: 3シート (Sheet1, Sheet2, OnlyInFile1)")
    print("  - sample_file2.xlsx: 3シート (Sheet1, Sheet2, OnlyInFile2)")
    print("\n予想される結果:")
    print("  - Sheet1: ファイル1に2行、ファイル2に2行の差分")
    print("  - Sheet2: 完全に異なる内容")
    print("  - OnlyInFile1: ファイル1にのみ存在（2行）")
    print("  - OnlyInFile2: ファイル2にのみ存在（2行）")


def test_excel_row_diff():
    """
    Excel Row Diff を実行してテスト
    """
    print("\n" + "="*50)
    print("Excel Row Diff を実行中...")
    print("="*50 + "\n")
    
    # excel_row_diff.py をインポートして実行
    try:
        from excel_row_diff import compare_excel_files
        compare_excel_files("sample_file1.xlsx", "sample_file2.xlsx", "test_result.xlsx")
        
        print("\n" + "="*50)
        print("テスト完了!")
        print("="*50)
        print("\n結果ファイル: test_result.xlsx")
        print("  - summary シートで差分の概要を確認できます")
        print("  - 各シートの差分が個別のシートに出力されています")
        
        return True
    except Exception as e:
        print(f"\nエラー: {e}")
        import traceback
        traceback.print_exc()
        return False


def cleanup():
    """
    テストファイルのクリーンアップ（オプション）
    """
    files = ["sample_file1.xlsx", "sample_file2.xlsx", "test_result.xlsx"]
    print("\nクリーンアップしますか？ (テストファイルを削除)")
    print("削除するファイル:", ", ".join(files))
    response = input("削除する場合は 'yes' と入力: ")
    
    if response.lower() == "yes":
        for file in files:
            if os.path.exists(file):
                os.remove(file)
                print(f"  ✓ {file} を削除")
        print("クリーンアップ完了")
    else:
        print("ファイルは保持されます")


if __name__ == "__main__":
    # サンプルファイルを作成
    create_sample_files()
    
    # テストを実行
    success = test_excel_row_diff()
    
    if success:
        print("\n✅ すべてのテストが成功しました！")
        
        # クリーンアップオプション
        if len(sys.argv) > 1 and sys.argv[1] == "--no-cleanup":
            print("\nテストファイルを保持します")
        else:
            cleanup()
    else:
        print("\n❌ テストが失敗しました")
        sys.exit(1)
