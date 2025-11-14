#!/usr/bin/env python3
"""
Excel Row Diff Tool

二つの Excel ファイルを比較し、同名シート同士の差分（片方にしか無い行）を抽出するツールです。
"""

import argparse
import sys
from pathlib import Path
from typing import Dict, List, Set, Tuple
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter


def get_row_as_tuple(sheet, row_idx: int) -> Tuple:
    """
    シートの指定行をタプルとして取得（比較用）
    """
    row = sheet[row_idx]
    return tuple(cell.value for cell in row)


def get_all_rows_as_set(sheet) -> Set[Tuple]:
    """
    シートの全行をセットとして取得（ヘッダー行を除く）
    """
    rows = set()
    # 2行目から取得（1行目はヘッダーと仮定）
    for row_idx in range(2, sheet.max_row + 1):
        row_data = get_row_as_tuple(sheet, row_idx)
        # 空行をスキップ
        if any(cell is not None and str(cell).strip() != "" for cell in row_data):
            rows.add(row_data)
    return rows


def copy_header_row(source_sheet, target_sheet):
    """
    ヘッダー行（1行目）をコピー
    """
    if source_sheet.max_row >= 1:
        for col_idx in range(1, source_sheet.max_column + 1):
            cell_value = source_sheet.cell(1, col_idx).value
            target_sheet.cell(1, col_idx, cell_value)


def write_rows_to_sheet(rows: List[Tuple], sheet, start_row: int = 2):
    """
    行データをシートに書き込み
    """
    for row_idx, row_data in enumerate(rows, start=start_row):
        for col_idx, cell_value in enumerate(row_data, start=1):
            sheet.cell(row_idx, col_idx, cell_value)


def copy_entire_sheet(source_sheet, target_sheet):
    """
    シート全体をコピー（すべての行とヘッダー）
    """
    for row_idx in range(1, source_sheet.max_row + 1):
        for col_idx in range(1, source_sheet.max_column + 1):
            cell_value = source_sheet.cell(row_idx, col_idx).value
            target_sheet.cell(row_idx, col_idx, cell_value)


def compare_excel_files(file1_path: str, file2_path: str, output_path: str):
    """
    2つのExcelファイルを比較し、差分を抽出
    
    Args:
        file1_path: 1つ目のExcelファイルのパス
        file2_path: 2つ目のExcelファイルのパス
        output_path: 出力Excelファイルのパス
    """
    print(f"Loading {file1_path}...")
    wb1 = load_workbook(file1_path, data_only=True)
    
    print(f"Loading {file2_path}...")
    wb2 = load_workbook(file2_path, data_only=True)
    
    # シート名を取得
    sheets1 = set(wb1.sheetnames)
    sheets2 = set(wb2.sheetnames)
    
    common_sheets = sheets1 & sheets2
    only_in_file1 = sheets1 - sheets2
    only_in_file2 = sheets2 - sheets1
    
    print(f"\n共通シート: {len(common_sheets)}")
    print(f"ファイル1のみ: {len(only_in_file1)}")
    print(f"ファイル2のみ: {len(only_in_file2)}")
    
    # 出力用ワークブックを作成
    output_wb = Workbook()
    # デフォルトシートを削除
    if "Sheet" in output_wb.sheetnames:
        output_wb.remove(output_wb["Sheet"])
    
    # サマリー情報を格納
    summary_data = []
    
    # 共通シートの比較
    for sheet_name in sorted(common_sheets):
        print(f"\n処理中: {sheet_name}")
        
        sheet1 = wb1[sheet_name]
        sheet2 = wb2[sheet_name]
        
        # 各シートの行を取得
        rows1 = get_all_rows_as_set(sheet1)
        rows2 = get_all_rows_as_set(sheet2)
        
        # 差分を計算
        only_in_1 = rows1 - rows2
        only_in_2 = rows2 - rows1
        
        print(f"  ファイル1のみ: {len(only_in_1)} 行")
        print(f"  ファイル2のみ: {len(only_in_2)} 行")
        
        # サマリーに追加
        summary_data.append({
            "シート名": sheet_name,
            "状態": "両方に存在",
            "ファイル1のみ": len(only_in_1),
            "ファイル2のみ": len(only_in_2),
            "総差分": len(only_in_1) + len(only_in_2)
        })
        
        # ファイル1のみの行を出力
        if only_in_1:
            sheet_name_safe = sheet_name.replace("/", "_").replace("\\", "_")[:31]
            output_sheet_name = f"{sheet_name_safe}_only1"
            output_sheet = output_wb.create_sheet(output_sheet_name)
            copy_header_row(sheet1, output_sheet)
            write_rows_to_sheet(sorted(only_in_1), output_sheet)
        
        # ファイル2のみの行を出力
        if only_in_2:
            sheet_name_safe = sheet_name.replace("/", "_").replace("\\", "_")[:31]
            output_sheet_name = f"{sheet_name_safe}_only2"
            output_sheet = output_wb.create_sheet(output_sheet_name)
            copy_header_row(sheet2, output_sheet)
            write_rows_to_sheet(sorted(only_in_2), output_sheet)
    
    # ファイル1にのみ存在するシート
    for sheet_name in sorted(only_in_file1):
        print(f"\n処理中: {sheet_name} (ファイル1のみ)")
        sheet1 = wb1[sheet_name]
        sheet_name_safe = sheet_name.replace("/", "_").replace("\\", "_")[:31]
        output_sheet_name = f"{sheet_name_safe}_only1"
        output_sheet = output_wb.create_sheet(output_sheet_name)
        copy_entire_sheet(sheet1, output_sheet)
        
        row_count = sheet1.max_row - 1 if sheet1.max_row > 1 else 0
        summary_data.append({
            "シート名": sheet_name,
            "状態": "ファイル1のみ",
            "ファイル1のみ": row_count,
            "ファイル2のみ": 0,
            "総差分": row_count
        })
    
    # ファイル2にのみ存在するシート
    for sheet_name in sorted(only_in_file2):
        print(f"\n処理中: {sheet_name} (ファイル2のみ)")
        sheet2 = wb2[sheet_name]
        sheet_name_safe = sheet_name.replace("/", "_").replace("\\", "_")[:31]
        output_sheet_name = f"{sheet_name_safe}_only2"
        output_sheet = output_wb.create_sheet(output_sheet_name)
        copy_entire_sheet(sheet2, output_sheet)
        
        row_count = sheet2.max_row - 1 if sheet2.max_row > 1 else 0
        summary_data.append({
            "シート名": sheet_name,
            "状態": "ファイル2のみ",
            "ファイル1のみ": 0,
            "ファイル2のみ": row_count,
            "総差分": row_count
        })
    
    # サマリーシートを作成（先頭に配置）
    if summary_data:
        summary_sheet = output_wb.create_sheet("summary", 0)
        
        # ヘッダー
        headers = ["シート名", "状態", "ファイル1のみ", "ファイル2のみ", "総差分"]
        for col_idx, header in enumerate(headers, start=1):
            summary_sheet.cell(1, col_idx, header)
        
        # データ
        for row_idx, data in enumerate(summary_data, start=2):
            summary_sheet.cell(row_idx, 1, data["シート名"])
            summary_sheet.cell(row_idx, 2, data["状態"])
            summary_sheet.cell(row_idx, 3, data["ファイル1のみ"])
            summary_sheet.cell(row_idx, 4, data["ファイル2のみ"])
            summary_sheet.cell(row_idx, 5, data["総差分"])
        
        # 合計行を追加
        total_row = len(summary_data) + 2
        summary_sheet.cell(total_row, 1, "合計")
        summary_sheet.cell(total_row, 3, sum(d["ファイル1のみ"] for d in summary_data))
        summary_sheet.cell(total_row, 4, sum(d["ファイル2のみ"] for d in summary_data))
        summary_sheet.cell(total_row, 5, sum(d["総差分"] for d in summary_data))
    
    # 出力ファイルを保存
    print(f"\n保存中: {output_path}")
    output_wb.save(output_path)
    print("完了!")
    
    # サマリーを表示
    print("\n=== サマリー ===")
    for data in summary_data:
        if data["総差分"] > 0:
            print(f"{data['シート名']}: {data['総差分']} 件の差分")


def main():
    """
    メイン関数
    """
    parser = argparse.ArgumentParser(
        description="Excel Row Diff - 2つのExcelファイルの行単位の差分を抽出",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
使用例:
  %(prog)s file1.xlsx file2.xlsx
  %(prog)s file1.xlsx file2.xlsx -o result.xlsx
        """
    )
    
    parser.add_argument(
        "file1",
        help="比較する1つ目のExcelファイル"
    )
    
    parser.add_argument(
        "file2",
        help="比較する2つ目のExcelファイル"
    )
    
    parser.add_argument(
        "-o", "--output",
        default="diff_result.xlsx",
        help="出力ファイル名 (デフォルト: diff_result.xlsx)"
    )
    
    args = parser.parse_args()
    
    # ファイルの存在確認
    file1_path = Path(args.file1)
    file2_path = Path(args.file2)
    
    if not file1_path.exists():
        print(f"エラー: ファイルが見つかりません: {args.file1}", file=sys.stderr)
        sys.exit(1)
    
    if not file2_path.exists():
        print(f"エラー: ファイルが見つかりません: {args.file2}", file=sys.stderr)
        sys.exit(1)
    
    try:
        compare_excel_files(str(file1_path), str(file2_path), args.output)
    except Exception as e:
        print(f"エラーが発生しました: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
